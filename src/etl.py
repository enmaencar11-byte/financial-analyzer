"""
ETL de FinAnalyzer Pro.
Extrae los datos de las plantillas .xlsx, transpone los años a filas
y consolida todos los emisores en un único CSV listo para el análisis.

Uso desde la raíz del proyecto:
    python -m src.etl
"""

from pathlib import Path

import openpyxl
import pandas as pd

from src.config import TODAS_LAS_COLUMNAS

# ── Rutas del proyecto ───────────────────────────────────────────
RAIZ = Path(__file__).resolve().parent.parent
DIR_ENTRADA = RAIZ / 'data' / 'raw'
DIR_SALIDA = RAIZ / 'data' / 'processed'
ARCHIVO_SALIDA = DIR_SALIDA / 'emisores_consolidado.csv'

# ── Columnas de los años en la plantilla ─────────────────────────
# B, C, D, E corresponden a los cuatro ejercicios.
COLUMNAS_ANIOS = [2, 3, 4, 5]
FILA_ENCABEZADO_ANIOS = 3

# ── Mapeo fila -> nombre de columna ──────────────────────────────
# Se mapea por NÚMERO DE FILA y no por texto de la etiqueta.
# Si algún día se corrige una tilde o se reformula una etiqueta en
# la plantilla, el ETL sigue funcionando. Si en cambio se INSERTA o
# ELIMINA una fila, estos números deben actualizarse.
FILAS_BALANCE = {
    5:  'efectivo_equivalentes',
    6:  'cuentas_por_cobrar',
    7:  'inventarios',
    8:  'otros_activos_corrientes',
    9:  'activo_corriente_total',
    11: 'propiedad_planta_equipo_neto',
    12: 'activos_intangibles',
    13: 'otros_activos_no_corrientes',
    14: 'activo_no_corriente_total',
    15: 'activo_total',
    17: 'cuentas_por_pagar',
    18: 'deuda_financiera_corto_plazo',
    19: 'otros_pasivos_corrientes',
    20: 'pasivo_corriente_total',
    22: 'deuda_financiera_largo_plazo',
    23: 'otros_pasivos_no_corrientes',
    24: 'pasivo_no_corriente_total',
    25: 'pasivo_total',
    27: 'capital_social',
    28: 'utilidades_retenidas',
    29: 'otras_reservas',
    30: 'patrimonio_controladora',
    31: 'participaciones_no_controladoras',
    32: 'patrimonio_total',
    38: 'deuda_financiera_total',
    39: 'capital_trabajo_neto',
}

FILAS_RESULTADOS = {
    5:  'ventas_netas',
    7:  'costo_ventas',
    8:  'utilidad_bruta',
    10: 'gastos_ventas',
    11: 'gastos_administracion',
    12: 'depreciacion_amortizacion',
    13: 'otros_gastos_operativos',
    14: 'total_gastos_operativos',
    15: 'utilidad_operativa',
    17: 'ingresos_financieros',
    18: 'gastos_financieros',
    19: 'otros_ingresos_netos',
    20: 'utilidad_antes_impuestos',
    22: 'impuesto_renta',
    23: 'utilidad_neta',
    25: 'utilidad_neta_controladora',
    26: 'utilidad_neta_no_controladoras',
    29: 'ebitda',
    30: 'tasa_impuesto_efectiva',
    31: 'nopat',
}

# Fila del activo total. Decide si una columna de ejercicio trae datos:
# una empresa sin activo no existe, de modo que un cero ahí indica
# columna vacía y no una cifra.
FILA_ACTIVO_TOTAL = 15

FILA_CUADRE = 35          # Balance General: debe ser 0
FILA_DISTRIBUCION = 27    # Estado de Resultados: debe ser 0
TOLERANCIA_CUADRE = 1.0   # margen por redondeo en la unidad declarada


# Hoja 'Datos Generales'
FILA_NOMBRE_REAL = 4
FILA_NOMBRE_TRABAJO = 5


def _valor(hoja, fila, columna):
    """Lee una celda y devuelve 0.0 si está vacía."""
    v = hoja.cell(row=fila, column=columna).value
    if v is None or v == '':
        return 0.0
    return float(v)


def _nombre_emisor(libro):
    """
    Obtiene el nombre del emisor desde la hoja 'Datos Generales'.
    Da prioridad al nombre de trabajo (anonimizado) si fue declarado.
    """
    hoja = libro['Datos Generales']
    nombre_trabajo = hoja.cell(row=FILA_NOMBRE_TRABAJO, column=2).value
    nombre_real = hoja.cell(row=FILA_NOMBRE_REAL, column=2).value
    nombre = nombre_trabajo or nombre_real
    if not nombre:
        raise ValueError("La hoja 'Datos Generales' no tiene nombre de emisor.")
    return str(nombre).strip()


def _ejercicios(hoja_balance, emisor=''):
    """
    Devuelve los ejercicios que la plantilla trae efectivamente llenos,
    como pares (columna, año).

    Una plantilla se llena por partes: se transcribe un ejercicio, se
    comprueba, y después los demás. Recorrer siempre las cuatro columnas
    tenía dos consecuencias, ambas indeseables. Una columna sin año en el
    encabezado interrumpía la lectura al intentar convertir un valor
    ausente. Y una columna con año pero sin cifras producía un ejercicio
    entero en ceros, que es peor que el error: no interrumpe nada y
    contamina los indicadores con un emisor sin activo, sin ventas y con
    divisiones por cero.

    Se considera lleno el ejercicio que declara año y presenta activo
    total distinto de cero.
    """
    ejercicios = []
    for col in COLUMNAS_ANIOS:
        anio = hoja_balance.cell(row=FILA_ENCABEZADO_ANIOS, column=col).value
        if anio is None or str(anio).strip() == '':
            continue
        if abs(_valor(hoja_balance, FILA_ACTIVO_TOTAL, col)) < 1e-9:
            continue
        ejercicios.append((col, int(anio)))

    if not ejercicios:
        raise ValueError(
            f'{emisor or "La plantilla"} no trae ningún ejercicio con datos.\n'
            f'Compruebe que la fila {FILA_ACTIVO_TOTAL} del Balance General '
            f'tenga el activo total, y que el archivo se haya abierto y '
            f'guardado en Excel al menos una vez después de llenarlo: las '
            f'celdas con fórmula se leen vacías en caso contrario.')
    return ejercicios


def _verificar_cuadre(hoja_balance, emisor, ejercicios):
    """
    Comprueba que activo = pasivo + patrimonio en cada ejercicio.
    Un descuadre indica error de transcripción, no error del emisor.
    """
    errores = []
    for col, anio in ejercicios:
        cuadre = _valor(hoja_balance, FILA_CUADRE, col)
        if abs(cuadre) > TOLERANCIA_CUADRE:
            errores.append(f'  {emisor} {anio}: descuadre de {cuadre:,.2f}')
    return errores

def _verificar_distribucion(hoja_resultados, emisor, ejercicios):
    """
    Comprueba que la utilidad neta sea igual a la suma de la porción
    atribuible a la controladora y la atribuible a participaciones no
    controladoras. En estados individuales la segunda es cero.
    """
    errores = []
    for col, anio in ejercicios:
        dif = _valor(hoja_resultados, FILA_DISTRIBUCION, col)
        if abs(dif) > TOLERANCIA_CUADRE:
            errores.append(f'  {emisor} {anio}: la distribución del resultado difiere en {dif:,.2f}')
    return errores


def leer_plantilla(ruta):
    """
    Lee una plantilla y devuelve un DataFrame con una fila por ejercicio.

    data_only=True devuelve el ÚLTIMO VALOR CALCULADO por Excel, no la
    fórmula. Esto exige que el archivo haya sido abierto y guardado en
    Excel al menos una vez después de llenarlo; de lo contrario, todas
    las celdas con fórmula se leerán como vacías.
    """
    libro = openpyxl.load_workbook(ruta, data_only=True)

    faltantes = [h for h in ('Datos Generales', 'Balance General',
                             'Estado de Resultados') if h not in libro.sheetnames]
    if faltantes:
        raise ValueError(f'{ruta.name}: faltan las hojas {faltantes}')

    balance = libro['Balance General']
    resultados = libro['Estado de Resultados']

    emisor = _nombre_emisor(libro)
    ejercicios = _ejercicios(balance, emisor)

    errores = _verificar_cuadre(balance, emisor, ejercicios)
    if errores:
        raise ValueError('Descuadre en el balance:\n' + '\n'.join(errores))

    errores = _verificar_distribucion(resultados, emisor, ejercicios)
    if errores:
        raise ValueError('Distribución del resultado inconsistente:\n' + '\n'.join(errores))
 

    filas = []
    for col, anio in ejercicios:
        registro = {'empresa': emisor, 'anio': anio}
        for fila, columna in FILAS_BALANCE.items():
            registro[columna] = _valor(balance, fila, col)
        for fila, columna in FILAS_RESULTADOS.items():
            registro[columna] = _valor(resultados, fila, col)
        filas.append(registro)

    return pd.DataFrame(filas)


def consolidar(dir_entrada=DIR_ENTRADA, archivo_salida=ARCHIVO_SALIDA):
    """Procesa todas las plantillas y escribe el CSV consolidado."""
    plantillas = sorted(p for p in dir_entrada.glob('*.xlsx')
                        if not p.name.startswith('~$'))

    if not plantillas:
        raise FileNotFoundError(f'No hay archivos .xlsx en {dir_entrada}')

    marcos = []
    for ruta in plantillas:
        print(f'  Leyendo {ruta.name} ...', end=' ')
        df = leer_plantilla(ruta)
        print(f'{len(df)} ejercicios de {df["empresa"].iloc[0]}')
        marcos.append(df)

    consolidado = pd.concat(marcos, ignore_index=True)

    ausentes = [c for c in TODAS_LAS_COLUMNAS if c not in consolidado.columns]
    if ausentes:
        raise ValueError(f'Faltan columnas exigidas por config.py: {ausentes}')

    # Se ordenan las columnas según config.py para que el CSV tenga
    # siempre la misma estructura, sin importar el orden de lectura.
    consolidado = consolidado[TODAS_LAS_COLUMNAS]
    consolidado = consolidado.sort_values(['empresa', 'anio']).reset_index(drop=True)

    archivo_salida.parent.mkdir(parents=True, exist_ok=True)
    consolidado.to_csv(archivo_salida, index=False, encoding='utf-8-sig')

    return consolidado


if __name__ == '__main__':
    print(f'\nLeyendo plantillas desde {DIR_ENTRADA}\n')
    datos = consolidar()
    print(f'\nConsolidado: {len(datos)} filas x {len(datos.columns)} columnas')
    print(f'Emisores: {", ".join(sorted(datos["empresa"].unique()))}')
    print(f'Ejercicios: {datos["anio"].min()} a {datos["anio"].max()}')
    print(f'Guardado en {ARCHIVO_SALIDA}\n')